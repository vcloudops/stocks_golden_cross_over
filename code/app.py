import os
from datetime import datetime
import pandas as pd
import yfinance as yf
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import PatternFill

# -----------------------------
# Configuration
# -----------------------------
NIFTY_50_TICKERS = [
    'RELIANCE.NS', 'HDFCBANK.NS', 'ICICIBANK.NS', 'TCS.NS', 'INFY.NS',
    'KOTAKBANK.NS', 'HINDUNILVR.NS', 'SBIN.NS', 'BAJFINANCE.NS', 'AXISBANK.NS'
    # Add more tickers as needed
]

OUTPUT_DIR = "./output"
SMA_SHORT = 50
SMA_LONG = 200
RECENT_DAYS = 30  # highlight recent crosses

os.makedirs(OUTPUT_DIR, exist_ok=True)

# -----------------------------
# Helper function
# -----------------------------
def check_golden_cross(ticker):
    try:
        data = yf.download(ticker, period="1y", auto_adjust=True)
        if data.empty or len(data) < SMA_LONG:
            print(f"Skipping {ticker}: Not enough data")
            return None

        data['SMA50'] = data['Close'].rolling(SMA_SHORT).mean()
        data['SMA200'] = data['Close'].rolling(SMA_LONG).mean()

        cross = (data['SMA50'] > data['SMA200']) & (data['SMA50'].shift(1) <= data['SMA200'].shift(1))
        cross_dates = data.index[cross]

        if not cross_dates.empty:
            last_cross = cross_dates[-1]
            # Handle tz-aware datetime
            last_cross_naive = last_cross.tz_localize(None)
            days_ago = (datetime.today() - last_cross_naive).days

            return {
                "Ticker": ticker,
                "Last Price": data['Close'].iloc[-1],
                "SMA50": data['SMA50'].iloc[-1],
                "SMA200": data['SMA200'].iloc[-1],
                "Golden Cross Date": last_cross_naive.date(),
                "Days Since Golden Cross": days_ago
            }
    except Exception as e:
        print(f"Error processing {ticker}: {e}")
    return None

# -----------------------------
# Collect data
# -----------------------------
results = []
for ticker in NIFTY_50_TICKERS:
    res = check_golden_cross(ticker)
    if res:
        results.append(res)

if not results:
    print("No stocks above the golden cross.")
    exit(0)

df = pd.DataFrame(results)
df.sort_values('Days Since Golden Cross', inplace=True)

# -----------------------------
# Save CSV
# -----------------------------
csv_file = os.path.join(OUTPUT_DIR, "nifty50_golden_cross.csv")
df.to_csv(csv_file, index=False)
print(f"CSV saved: {csv_file}")

# -----------------------------
# Save Excel with highlighting
# -----------------------------
excel_file = os.path.join(OUTPUT_DIR, "nifty50_golden_cross.xlsx")
with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
    df.to_excel(writer, index=False, sheet_name='GoldenCross')
    sheet = writer.sheets['GoldenCross']
    for idx, days in enumerate(df['Days Since Golden Cross'], start=2):
        if days <= RECENT_DAYS:
            for col in range(1, len(df.columns)+1):
                cell = sheet.cell(row=idx, column=col)
                cell.fill = PatternFill(start_color="C6EFCE", fill_type="solid")
print(f"Excel saved: {excel_file}")

# -----------------------------
# Generate PDF report
# -----------------------------
pdf_file = os.path.join(OUTPUT_DIR, "nifty50_golden_cross_report.pdf")
with PdfPages(pdf_file) as pdf:
    for ticker in df['Ticker']:
        data = yf.download(ticker, period="1y", auto_adjust=True)
        data['SMA50'] = data['Close'].rolling(SMA_SHORT).mean()
        data['SMA200'] = data['Close'].rolling(SMA_LONG).mean()

        plt.figure(figsize=(10,6))
        plt.plot(data['Close'], label='Close', color='blue')
        plt.plot(data['SMA50'], label='SMA50', color='green')
        plt.plot(data['SMA200'], label='SMA200', color='red')

        cross = (data['SMA50'] > data['SMA200']) & (data['SMA50'].shift(1) <= data['SMA200'].shift(1))
        cross_dates = data.index[cross]
        for date in cross_dates:
            plt.axvline(date, color='orange', linestyle='--', alpha=0.7)

        plt.title(f"{ticker} Golden Cross Chart")
        plt.xlabel("Date")
        plt.ylabel("Price")
        plt.legend()
        pdf.savefig()
        plt.close()

print(f"PDF report saved: {pdf_file}")